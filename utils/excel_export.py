import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def aplicar_estilo_excel(ws, titulo, columnas, df):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
    titulo_cell = ws.cell(row=1, column=1)
    titulo_cell.value = titulo.upper()
    titulo_cell.font = Font(name='Calibri', size=18, bold=True, color="FFFFFF")
    titulo_cell.alignment = Alignment(horizontal="center", vertical="center")
    titulo_cell.fill = PatternFill("solid", fgColor="7B2026")

    header_fill = PatternFill("solid", fgColor="2C5282")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for idx, col in enumerate(columnas, start=1):
        cell = ws.cell(row=2, column=idx, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for r_idx, row in enumerate(df.itertuples(index=False), start=3):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = Alignment(horizontal="left")
            cell.border = border

            if columnas[c_idx - 1].lower() == "última donación" and str(value).strip().lower() == "sin donacion":
                cell.font = Font(color="FF0000")

    # Ajuste automático de ancho con límites
    for i, col in enumerate(columnas, start=1):
        col_letter = get_column_letter(i)
        max_length = max(
            len(str(ws.cell(row=row, column=i).value or "")) for row in range(1, ws.max_row + 1)
        )
        if columnas[i - 1].lower() in ["id", "id donación"]:
            ws.column_dimensions[col_letter].width = 6
        else:
            ws.column_dimensions[col_letter].width = min(max(max_length + 2, 10), 40)

def exportar_donantes_excel(nombre_archivo, donantes):
    columnas = ["ID", "Nombre", "Teléfono", "Email", "Cédula", "Dirección", "Última Donación"]
    df = pd.DataFrame([fila[:len(columnas)] for fila in donantes], columns=columnas)
    wb = Workbook()
    ws = wb.active
    ws.title = "Donantes"
    aplicar_estilo_excel(ws, "Base de Datos de Donantes", columnas, df)
    wb.save(nombre_archivo)

def exportar_donaciones_excel(nombre_archivo, donaciones):
    columnas = ["ID Donación", "Fecha", "Descripción", "Nombre", "Cédula", "Teléfono", "Email"]
    filas = []

    for d in donaciones:
        fila = [
            d.get("ID_DONACION"),
            d.get("FECHA"),
            d.get("DESCRI"),
            d.get("NOMBRE"),
            d.get("CEDULA"),
            d.get("TELEFONO"),
            d.get("CORREO")
        ]
        filas.append(fila)

    df = pd.DataFrame(filas, columns=columnas)
    wb = Workbook()
    ws = wb.active
    ws.title = "Donaciones"
    aplicar_estilo_excel(ws, "Base de Datos de Donaciones", columnas, df)
    wb.save(nombre_archivo)

def exportar_voluntarios_excel(nombre_archivo, voluntarios):
    columnas = ["ID", "Nombre", "Cédula", "Teléfono", "Email", "Disponibilidad"]
    df = pd.DataFrame([fila[:len(columnas)] for fila in voluntarios], columns=columnas)
    wb = Workbook()
    ws = wb.active
    ws.title = "Voluntarios"
    aplicar_estilo_excel(ws, "Base de Datos de Voluntarios", columnas, df)
    wb.save(nombre_archivo)

def exportar_jornadas_excel(nombre_archivo, jornadas):
    columnas = ["ID", "Fecha", "Descripción", "Ubicación", "Estatus"]
    df = pd.DataFrame([fila[:len(columnas)] for fila in jornadas], columns=columnas)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jornadas"
    aplicar_estilo_excel(ws, "Base de Datos de Jornadas", columnas, df)
    wb.save(nombre_archivo)
